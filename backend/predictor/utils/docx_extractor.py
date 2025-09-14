"""DOCX text extraction utilities for syllabus processing."""

import os
import logging
from typing import Optional, Dict, Any, List
from pathlib import Path

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    Document = None

logger = logging.getLogger(__name__)


class DOCXExtractor:
    """Extract text content from DOCX files."""

    def __init__(self):
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for DOCX processing. Install with: pip install python-docx")

    def extract_text(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text from DOCX file.

        Args:
            file_path: Path to the DOCX file

        Returns:
            Dict containing extracted text and metadata
        """
        try:
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"DOCX file not found: {file_path}")

            doc = Document(file_path)

            # Extract metadata
            metadata = {
                'title': getattr(doc.core_properties, 'title', ''),
                'author': getattr(doc.core_properties, 'author', ''),
                'subject': getattr(doc.core_properties, 'subject', ''),
                'created': getattr(doc.core_properties, 'created', ''),
                'modified': getattr(doc.core_properties, 'modified', ''),
            }

            # Extract text from paragraphs
            paragraphs = []
            for para_num, paragraph in enumerate(doc.paragraphs):
                text = paragraph.text.strip()
                if text:  # Only add non-empty paragraphs
                    paragraphs.append({
                        'paragraph_number': para_num + 1,
                        'content': text,
                        'char_count': len(text)
                    })

            # Extract text from tables
            tables_content = []
            for table_num, table in enumerate(doc.tables):
                table_text = []
                for row_num, row in enumerate(table.rows):
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        table_text.append(' | '.join(row_text))

                if table_text:
                    tables_content.append({
                        'table_number': table_num + 1,
                        'content': '\n'.join(table_text),
                        'rows': len(table.rows)
                    })

            # Combine all text
            all_text_parts = []

            # Add paragraph text
            para_text = '\n\n'.join([p['content'] for p in paragraphs])
            if para_text:
                all_text_parts.append(para_text)

            # Add table text
            for table in tables_content:
                all_text_parts.append(f"\nTable {table['table_number']}:\n{table['content']}")

            full_text = '\n\n'.join(all_text_parts)

            result = {
                'success': True,
                'text': full_text,
                'paragraphs': paragraphs,
                'tables': tables_content,
                'metadata': metadata,
                'total_paragraphs': len(paragraphs),
                'total_tables': len(tables_content),
                'total_chars': len(full_text)
            }

            logger.info(f"Successfully extracted text from DOCX: {len(full_text)} characters, {len(paragraphs)} paragraphs, {len(tables_content)} tables")
            return result

        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {e}")
            return {
                'success': False,
                'error': str(e),
                'text': '',
                'paragraphs': [],
                'tables': [],
                'metadata': {},
                'total_paragraphs': 0,
                'total_tables': 0,
                'total_chars': 0
            }

    def get_paragraph_count(self, file_path: str) -> int:
        """Get the number of paragraphs in a DOCX file."""
        try:
            doc = Document(file_path)
            return len([p for p in doc.paragraphs if p.text.strip()])
        except Exception as e:
            logger.error(f"Error getting paragraph count for DOCX {file_path}: {e}")
            return 0

    def validate_docx(self, file_path: str) -> bool:
        """Validate if the file is a valid DOCX."""
        try:
            Document(file_path)
            return True
        except Exception as e:
            logger.error(f"Invalid DOCX file {file_path}: {e}")
            return False
